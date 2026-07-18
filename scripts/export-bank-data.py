"""Exports the "27 banks" .xlsm workbook into the static JSON files the
"Cơ bản" (Fundamentals) section reads at client/public/data/banks/.

Usage:
    python3 scripts/export-bank-data.py path/to/File_27_Ngan_hang.xlsm

Requires openpyxl (`pip install openpyxl`).

Row numbers below are hardcoded against this workbook's specific layout
(one uniform 341-row template shared by all 27 bank sheets — verified
identical row-for-row across banks) and its "Công thức" sheet's per-metric
formulas (column K, "Theo dòng"). Every ratio here was cross-checked against
that sheet's own live "Chi tiết" dashboard values for NAB Q1 2026 before
being trusted — see the git history for that verification. If a future
version of the workbook reorders rows or adds/removes bank sheets, this
script needs updating to match; it does not read the formulas dynamically.

To refresh the site's data after the user sends an updated workbook: run
this script pointing at the new file, then check in the changed files
under client/public/data/banks/.
"""

import json
import math
import os
import sys

import openpyxl
from openpyxl.utils import column_index_from_string

if len(sys.argv) < 2:
    print("Usage: python3 export-bank-data.py <path-to-xlsm>")
    sys.exit(1)

SRC = sys.argv[1]
REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(REPO_ROOT, "client", "public", "data", "banks")

os.makedirs(OUT_DIR, exist_ok=True)
os.makedirs(os.path.join(OUT_DIR, "statements"), exist_ok=True)

wb = openpyxl.load_workbook(SRC, data_only=True, keep_vba=True)

# ---- Bank metadata from "Công thức" sheet rows 2-28 ----
formula_ws = wb["Công thức"]
banks = []
for r in range(2, 29):
    symbol = formula_ws.cell(row=r, column=2).value
    if not symbol:
        continue
    banks.append(
        {
            "symbol": symbol,
            "name": formula_ws.cell(row=r, column=3).value,
            "group": formula_ws.cell(row=r, column=4).value or "",
            "exchange": formula_ws.cell(row=r, column=5).value or "",
        }
    )
print(f"{len(banks)} banks:", [b["symbol"] for b in banks])

YEAR_COLS = [column_index_from_string(c) for c in ["B", "C", "D", "E", "F", "G", "H", "I"]]
YEAR_LABELS = ["2018", "2019", "2020", "2021", "2022", "2023", "2024", "2025"]

QUARTER_COLS = list(range(column_index_from_string("K"), column_index_from_string("AQ") + 1))


def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
            return None
        return v
    return None


def safe_div(a, b):
    if a is None or b is None or b == 0:
        return None
    return a / b


class BankSheet:
    def __init__(self, ws):
        self.ws = ws
        self.labels = [ws.cell(row=r, column=1).value for r in range(1, 342)]

    def row(self, r, col_idx):
        return num(self.ws.cell(row=r, column=col_idx).value)


def build_period_series(sheet: BankSheet, col_indices, labels, is_quarterly: bool):
    """Compute raw rows + derived ratio metrics for one granularity (quarter or year)."""
    n = len(col_indices)

    def r(row_num, i):
        if i < 0 or i >= n:
            return None
        return sheet.row(row_num, col_indices[i])

    def r_sum4(row_num, i):
        """Trailing 4-period sum ending at i (quarterly annualization)."""
        if i - 3 < 0:
            return None
        vals = [r(row_num, k) for k in range(i - 3, i + 1)]
        if any(v is None for v in vals):
            return None
        return sum(vals)

    metrics = {
        "totalAssets": [],
        "loans": [],
        "deposits": [],
        "equity": [],
        "netProfit": [],
        "nii": [],
        "toi": [],
        "operatingExpense": [],
        "casa": [],
        "ldr": [],
        "llr": [],
        "cir": [],
        "npl": [],
        "roe": [],
        "roa": [],
        "roe4q": [],
        "roa4q": [],
        "nim": [],
        "yoea": [],
        "cof": [],
        "creditGrowthQoQ": [],
        "depositGrowthQoQ": [],
        "car": [],
    }

    for i in range(n):
        totalAssets = r(3, i)
        loans = r(14, i)
        deposits = r(57, i)
        equity = r(66, i)
        netProfit = r(115, i)
        nii = r(94, i)
        toi = r(107, i)
        opex = r(108, i)
        car = r(341, i)

        metrics["totalAssets"].append(totalAssets)
        metrics["loans"].append(loans)
        metrics["deposits"].append(deposits)
        metrics["equity"].append(equity)
        metrics["netProfit"].append(netProfit)
        metrics["nii"].append(nii)
        metrics["toi"].append(toi)
        metrics["operatingExpense"].append(-opex if opex is not None else None)
        metrics["car"].append(car)

        # CASA = tiền gửi không kỳ hạn (243) / tiền gửi khách hàng (57)
        metrics["casa"].append(safe_div(r(243, i), deposits))

        # LDR = cho vay KH (14) / (tiền gửi TCTD(55) + tiền gửi KH(57) + phát hành GTCG(60))
        funding = None
        f55, f57, f60 = r(55, i), deposits, r(60, i)
        if f55 is not None and f57 is not None and f60 is not None:
            funding = f55 + f57 + f60
        metrics["ldr"].append(safe_div(loans, funding))

        # LLR = -dự phòng rủi ro cho vay KH(16) / (nợ dưới TC(190)+nghi ngờ(191)+mất vốn(192))
        r190, r191, r192 = r(190, i), r(191, i), r(192, i)
        bad_debt = None
        if r190 is not None and r191 is not None and r192 is not None:
            bad_debt = r190 + r191 + r192
        r16 = r(16, i)
        llr = safe_div(-r16 if r16 is not None else None, bad_debt)
        metrics["llr"].append(llr)

        # CIR = -chi phí QLDN(108) / TOI(107)
        cir = safe_div(-opex if opex is not None else None, toi)
        metrics["cir"].append(cir)

        # NPL = bad_debt / cho vay KH gộp (15)
        metrics["npl"].append(safe_div(bad_debt, r(15, i)))

        # ROE/ROA — single period (non-annualized), avg balance vs 1 period back
        equity_prev1 = r(66, i - 1) if i - 1 >= 0 else None
        assets_prev1 = r(3, i - 1) if i - 1 >= 0 else None
        equity_avg1 = (equity + equity_prev1) / 2.0 if equity is not None and equity_prev1 is not None else None
        assets_avg1 = (
            (totalAssets + assets_prev1) / 2.0 if totalAssets is not None and assets_prev1 is not None else None
        )
        metrics["roe"].append(safe_div(netProfit, equity_avg1))
        metrics["roa"].append(safe_div(netProfit, assets_avg1))

        earning_assets = None
        e6, e10, e14, e20 = r(6, i), r(10, i), r(14, i), r(20, i)
        if None not in (e6, e10, e14, e20):
            earning_assets = e6 + e10 + e14 + e20
        earning_assets_prev1 = None
        if i - 1 >= 0:
            p6, p10, p14, p20 = r(6, i - 1), r(10, i - 1), r(14, i - 1), r(20, i - 1)
            if None not in (p6, p10, p14, p20):
                earning_assets_prev1 = p6 + p10 + p14 + p20
        earning_assets_avg1 = (
            (earning_assets + earning_assets_prev1) / 2.0
            if earning_assets is not None and earning_assets_prev1 is not None
            else None
        )

        # COF's denominator is row 54 ("Tiền gửi VÀ VAY các TCTD khác" — deposits +
        # borrowings) not row 55 (just the deposits sub-component) that LDR uses —
        # different row per the Công thức sheet, so this needs its own funding value
        # rather than reusing LDR's `funding`.
        r54 = r(54, i)
        funding_cof = r54 + deposits + f60 if None not in (r54, deposits, f60) else None
        funding_cof_prev1 = None
        if i - 1 >= 0:
            p54, p57, p60 = r(54, i - 1), r(57, i - 1), r(60, i - 1)
            if None not in (p54, p57, p60):
                funding_cof_prev1 = p54 + p57 + p60
        funding_cof_avg1 = (
            (funding_cof + funding_cof_prev1) / 2.0
            if funding_cof is not None and funding_cof_prev1 is not None
            else None
        )

        r95, r96 = r(95, i), r(96, i)

        if is_quarterly:
            # Trailing-4-quarter annualized versions
            equity_prev4 = r(66, i - 4) if i - 4 >= 0 else None
            assets_prev4 = r(3, i - 4) if i - 4 >= 0 else None
            netProfit_4q = r_sum4(115, i)
            equity_avg4 = (
                (equity + equity_prev4) / 2.0 if equity is not None and equity_prev4 is not None else None
            )
            assets_avg4 = (
                (totalAssets + assets_prev4) / 2.0
                if totalAssets is not None and assets_prev4 is not None
                else None
            )
            metrics["roe4q"].append(safe_div(netProfit_4q, equity_avg4))
            metrics["roa4q"].append(safe_div(netProfit_4q, assets_avg4))

            nii_4q = r_sum4(94, i)
            r95_4q = r_sum4(95, i)
            r96_4q = r_sum4(96, i)
            metrics["nim"].append(safe_div(nii_4q, earning_assets_avg1))
            metrics["yoea"].append(safe_div(r95_4q, earning_assets_avg1))
            metrics["cof"].append(safe_div(-r96_4q if r96_4q is not None else None, funding_cof_avg1))
        else:
            # Annual columns already represent the full year — use directly, no trailing sum.
            metrics["roe4q"].append(metrics["roe"][-1])
            metrics["roa4q"].append(metrics["roa"][-1])
            metrics["nim"].append(safe_div(nii, earning_assets_avg1))
            metrics["yoea"].append(safe_div(r95, earning_assets_avg1))
            metrics["cof"].append(safe_div(-r96 if r96 is not None else None, funding_cof_avg1))

        loans_prev1 = r(14, i - 1) if i - 1 >= 0 else None
        deposits_prev1 = r(57, i - 1) if i - 1 >= 0 else None
        metrics["creditGrowthQoQ"].append(
            safe_div(loans - loans_prev1 if loans is not None and loans_prev1 is not None else None, loans_prev1)
        )
        metrics["depositGrowthQoQ"].append(
            safe_div(
                deposits - deposits_prev1 if deposits is not None and deposits_prev1 is not None else None,
                deposits_prev1,
            )
        )

    return {"periods": labels, "metrics": metrics}


overview_rows = []

for b in banks:
    symbol = b["symbol"]
    if symbol not in wb.sheetnames:
        print("MISSING SHEET:", symbol)
        continue
    ws = wb[symbol]
    sheet = BankSheet(ws)

    quarter_labels = [ws.cell(row=2, column=c).value for c in QUARTER_COLS]
    quarterly = build_period_series(sheet, QUARTER_COLS, quarter_labels, is_quarterly=True)
    yearly = build_period_series(sheet, YEAR_COLS, YEAR_LABELS, is_quarterly=False)

    out = {
        "symbol": symbol,
        "name": b["name"],
        "group": b["group"],
        "exchange": b["exchange"],
        "quarter": quarterly,
        "year": yearly,
    }
    with open(os.path.join(OUT_DIR, f"{symbol}.json"), "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, separators=(",", ":"))

    # Raw statement export (341 rows), quarterly + yearly.
    raw_items = []
    for row_i in range(1, 342):
        label = sheet.labels[row_i - 1]
        if not label:
            continue
        raw_items.append(
            {
                "row": row_i,
                "name": label,
                "quarter": [sheet.row(row_i, c) for c in QUARTER_COLS],
                "year": [sheet.row(row_i, c) for c in YEAR_COLS],
            }
        )
    with open(os.path.join(OUT_DIR, "statements", f"{symbol}.json"), "w", encoding="utf-8") as f:
        json.dump(
            {"symbol": symbol, "quarterPeriods": quarter_labels, "yearPeriods": YEAR_LABELS, "items": raw_items},
            f,
            ensure_ascii=False,
            separators=(",", ":"),
        )

    # Latest-quarter snapshot row for the bank-comparison page.
    m = quarterly["metrics"]
    last = len(quarter_labels) - 1
    overview_rows.append(
        {
            "symbol": symbol,
            "name": b["name"],
            "group": b["group"],
            "exchange": b["exchange"],
            "period": quarter_labels[last],
            **{k: v[last] for k, v in m.items()},
        }
    )
    print(f"{symbol}: {len(quarter_labels)} quarters, {len(YEAR_LABELS)} years — done")

with open(os.path.join(OUT_DIR, "overview.json"), "w", encoding="utf-8") as f:
    json.dump(
        {"period": overview_rows[0]["period"] if overview_rows else None, "banks": overview_rows},
        f,
        ensure_ascii=False,
        separators=(",", ":"),
    )

with open(os.path.join(OUT_DIR, "meta.json"), "w", encoding="utf-8") as f:
    json.dump({"banks": banks}, f, ensure_ascii=False, indent=2)

print("DONE")
