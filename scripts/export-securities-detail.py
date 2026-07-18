"""Exports the securities-industry .xlsm workbook's per-company sheets and
"Dữ liệu Chung" table into the static JSON files the "Chi tiết mã chứng
khoán" dashboard reads at client/public/data/securities/detail/ —
companion to export-securities-data.py (14-core-ratio "Cơ bản" section)
and structurally identical to export-bank-detail.py, reverse-engineered
the same way from "Dữ liệu Biểu đồ Chi tiết"'s actual formulas (not
guessed), including two labeling bugs in that sheet caught by
cross-checking against the raw statement's own row labels:
  - Row 94 ("LN Môi giới") references cell B93 (this sheet's own
    "DT Môi giới") PLUS row 249, not row 249 alone — a same-sheet cell
    reference a naive Indexs()-only formula scan misses.
  - Rows 141/142 mislabel rows 221/222: the raw statement's own labels
    (verified directly) are row 220 = "Lãi bán tài sản", row 221 =
    "Chênh lệch đánh giá lại", row 222 = "Cổ tức & tiền lãi" — rows
    141/142's titles have 221 and 222 swapped.

Unlike the bank workbook, every series needed here is independently
derivable from the raw per-company rows at both quarterly and yearly
granularity — there's no "Dữ liệu Chung"-only metric requiring a
Q4-of-year snapshot fallback the way SMLR/sector-loan-% did for banks.
"""

import json
import math
import os
import sys
import openpyxl
from openpyxl.utils import column_index_from_string

SRC = sys.argv[1] if len(sys.argv) > 1 else "File_Chung_khoan.xlsm"
OUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "client/public/data/securities/detail")
os.makedirs(OUT_DIR, exist_ok=True)

wb = openpyxl.load_workbook(SRC, data_only=True, keep_vba=True)

formula_ws = wb["Công thức"]
COMPANIES = []
for r in range(2, 18):
    symbol = formula_ws.cell(row=r, column=2).value
    if symbol:
        COMPANIES.append(symbol)
print(f"{len(COMPANIES)} companies:", COMPANIES)

YEAR_COLS = [column_index_from_string(c) for c in ["B", "C", "D", "E", "F", "G", "H", "I"]]
YEAR_LABELS = ["2018", "2019", "2020", "2021", "2022", "2023", "2024", "2025"]
QUARTER_COLS = list(range(column_index_from_string("K"), column_index_from_string("AQ") + 1))


def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
            return None
        return v
    return None


def safe_div(a, b):
    if a is None or b is None or b == 0:
        return None
    return a / b


def growth(cur, prev):
    if cur is None or prev is None or prev == 0:
        return None
    return (cur - prev) / abs(prev)


def add(*vals):
    if all(v is None for v in vals):
        return None
    return sum(v or 0 for v in vals)


def sub(a, b):
    if a is None or b is None:
        return None
    return a - b


def shares(*series_list):
    """Per-index share of each series relative to their sum; an entire
    period is null (not just zeroed) if any component is missing."""
    n = len(series_list[0])
    out = [[] for _ in series_list]
    for i in range(n):
        vals = [s[i] for s in series_list]
        if any(v is None for v in vals):
            for o in out:
                o.append(None)
            continue
        total = sum(vals)
        for j, v in enumerate(vals):
            out[j].append(safe_div(v, total))
    return out


class Sheet:
    def __init__(self, ws):
        self.ws = ws

    def row(self, r, col_idx):
        return num(self.ws.cell(row=r, column=col_idx).value)


def compute_series(sheet: Sheet, cols: list, is_quarterly: bool) -> dict:
    n = len(cols)

    def r(row_num, i):
        return sheet.row(row_num, cols[i]) if 0 <= i < n else None

    def series(row_num):
        return [r(row_num, i) for i in range(n)]

    def avg1(row_num, i):
        cur, prev = r(row_num, i), r(row_num, i - 1) if i - 1 >= 0 else None
        return (cur + prev) / 2.0 if cur is not None and prev is not None else None

    s = {}

    # Revenue by segment
    dtFvtpl, dtHtm, dtAfs, dtChoVay, dtMoiGioi = series(219), series(223), series(225), series(224), series(227)
    dtHoatDong = series(237)
    dtKhac = [sub(t, add(a, b, c, d, e)) for t, a, b, c, d, e in zip(dtHoatDong, dtFvtpl, dtHtm, dtAfs, dtChoVay, dtMoiGioi)]
    s["dtFvtpl"], s["dtHtm"], s["dtAfs"], s["dtChoVay"], s["dtMoiGioi"], s["dtKhac"] = dtFvtpl, dtHtm, dtAfs, dtChoVay, dtMoiGioi, dtKhac
    s["dtHoatDong"] = dtHoatDong
    s["dtFvtplShare"], s["dtHtmShare"], s["dtAfsShare"], s["dtChoVayShare"], s["dtMoiGioiShare"], s["dtKhacShare"] = shares(
        dtFvtpl, dtHtm, dtAfs, dtChoVay, dtMoiGioi, dtKhac
    )
    s["revenueGrowth"] = [growth(dtHoatDong[i], dtHoatDong[i - 1]) if i >= 1 else None for i in range(n)]

    # Expense by segment. "CP mảng Cho vay" is not a direct row — it's an
    # implied funding-cost allocation proportional to lending's share of
    # total proprietary assets: -(interest income on loans) * loans /
    # (FVTPL+HTM+loans+AFS), per the validated chart-sheet formula.
    cpFvtpl, cpHtm, cpAfs, cpMoiGioi, cpKhac = series(239), series(243), series(245), series(249), series(255)
    f8, f9, f10, f11 = series(8), series(9), series(10), series(11)
    cpChoVay = []
    for i in range(n):
        denom = add(f8[i], f9[i], f10[i], f11[i])
        cpChoVay.append(-safe_div(dtChoVay[i], denom) * f10[i] if denom and dtChoVay[i] is not None and f10[i] is not None else None)

    # Profit by segment = revenue + expense (expense rows are already
    # negative in the raw statement, so this is a plain sum) — validated
    # against the chart-sheet's own formulas (e.g. row 98's B97 + ... ),
    # not the "Công thức" sheet's ambiguous +/- text.
    lng = series(257)
    lnFvtpl = [add(a, b) for a, b in zip(dtFvtpl, cpFvtpl)]
    lnHtm = [add(a, b) for a, b in zip(dtHtm, cpHtm)]
    lnAfs = [add(a, b) for a, b in zip(dtAfs, cpAfs)]
    lnChoVay = [add(a, b) for a, b in zip(dtChoVay, cpChoVay)]
    lnMoiGioi = [add(a, b) for a, b in zip(dtMoiGioi, cpMoiGioi)]
    lnKhac = [sub(t, add(a, b, c, d, e)) for t, a, b, c, d, e in zip(lng, lnFvtpl, lnHtm, lnAfs, lnChoVay, lnMoiGioi)]
    s["lnFvtpl"], s["lnHtm"], s["lnAfs"], s["lnChoVay"], s["lnMoiGioi"], s["lnKhac"] = (
        lnFvtpl, lnHtm, lnAfs, lnChoVay, lnMoiGioi, lnKhac,
    )
    s["lng"] = lng
    s["lnFvtplShare"], s["lnHtmShare"], s["lnAfsShare"], s["lnChoVayShare"], s["lnMoiGioiShare"], s["lnKhacShare"] = shares(
        lnFvtpl, lnHtm, lnAfs, lnChoVay, lnMoiGioi, lnKhac
    )
    s["blnFvtpl"] = [safe_div(a, b) for a, b in zip(lnFvtpl, dtFvtpl)]
    s["blnHtm"] = [safe_div(a, b) for a, b in zip(lnHtm, dtHtm)]
    s["blnAfs"] = [safe_div(a, b) for a, b in zip(lnAfs, dtAfs)]
    s["blnChoVay"] = [safe_div(a, b) for a, b in zip(lnChoVay, dtChoVay)]
    s["blnMoiGioi"] = [safe_div(a, b) for a, b in zip(lnMoiGioi, dtMoiGioi)]
    s["blnKhac"] = [safe_div(a, b) for a, b in zip(lnKhac, dtKhac)]

    # Pretax profit composition
    s["lnTuHDKD"] = series(272)
    s["lnTaiChinh"] = [add(a, b) for a, b in zip(series(258), series(263))]
    s["lnCtyLienKet"] = series(269)
    s["lntt"] = series(273)

    # FVTPL interest/gain composition — the raw statement's own row labels
    # (not "Dữ liệu Biểu đồ Chi tiết" rows 141/142's swapped titles): row
    # 220 = asset-sale gain, row 221 = revaluation, row 222 = dividend/
    # interest income, each paired with its FVTPL-cost-side counterpart.
    s["fvtplLaiBan"] = [add(a, b) for a, b in zip(series(220), series(240))]
    s["fvtplDanhGiaLai"] = [add(a, b) for a, b in zip(series(221), series(241))]
    s["fvtplCoTuc"] = [add(a, b) for a, b in zip(series(222), series(242))]
    s["fvtplLaiBanShare"], s["fvtplDanhGiaLaiShare"], s["fvtplCoTucShare"] = shares(
        s["fvtplLaiBan"], s["fvtplDanhGiaLai"], s["fvtplCoTuc"]
    )

    # Proprietary-book (tự doanh) assets
    tsFvtpl, tsHtm, tsAfs = f8, f9, f11
    s["tsFvtpl"], s["tsHtm"], s["tsAfs"] = tsFvtpl, tsHtm, tsAfs
    s["tsFvtplShare"], s["tsHtmShare"], s["tsAfsShare"] = shares(tsFvtpl, tsHtm, tsAfs)
    propTotal = [add(a, b, c) for a, b, c in zip(tsFvtpl, tsHtm, tsAfs)]
    s["propTotal"] = propTotal
    s["propAssetsGrowth"] = [growth(propTotal[i], propTotal[i - 1]) if i >= 1 else None for i in range(n)]

    # FVTPL trading yield: "Lãi gộp FVTPL" (= lnFvtpl, revenue+cost already
    # netted) over the average FVTPL book size across the period.
    fvtplAvgAssets = [
        (tsFvtpl[i] + tsFvtpl[i - 1]) / 2.0 if i >= 1 and None not in (tsFvtpl[i], tsFvtpl[i - 1]) else None
        for i in range(n)
    ]
    s["fvtplAvgAssets"] = fvtplAvgAssets
    s["fvtplYield"] = [safe_div(a, b) for a, b in zip(lnFvtpl, fvtplAvgAssets)]

    # Proprietary-book composition by instrument type, per book and
    # combined across all three books — both absolute and share.
    fvtplListed, fvtplUnlisted, fvtplFund, fvtplBond, fvtplMoneyMkt = series(578), series(579), series(580), series(581), series(583)
    htmListed, htmUnlisted, htmFund, htmBond, htmMoneyMkt = series(603), series(604), series(605), series(606), series(607)
    afsListed, afsUnlisted, afsFund, afsBond, afsMoneyMkt = series(591), series(592), series(593), series(594), series(595)
    s["fvtplListed"], s["fvtplUnlisted"], s["fvtplFund"], s["fvtplBond"], s["fvtplMoneyMkt"] = (
        fvtplListed, fvtplUnlisted, fvtplFund, fvtplBond, fvtplMoneyMkt,
    )
    s["htmListed"], s["htmUnlisted"], s["htmFund"], s["htmBond"], s["htmMoneyMkt"] = (
        htmListed, htmUnlisted, htmFund, htmBond, htmMoneyMkt,
    )
    s["afsListed"], s["afsUnlisted"], s["afsFund"], s["afsBond"], s["afsMoneyMkt"] = (
        afsListed, afsUnlisted, afsFund, afsBond, afsMoneyMkt,
    )
    s["fvtplListedShare"], s["fvtplUnlistedShare"], s["fvtplFundShare"], s["fvtplBondShare"], s["fvtplMoneyMktShare"] = shares(
        fvtplListed, fvtplUnlisted, fvtplFund, fvtplBond, fvtplMoneyMkt
    )
    s["htmListedShare"], s["htmUnlistedShare"], s["htmFundShare"], s["htmBondShare"], s["htmMoneyMktShare"] = shares(
        htmListed, htmUnlisted, htmFund, htmBond, htmMoneyMkt
    )
    s["afsListedShare"], s["afsUnlistedShare"], s["afsFundShare"], s["afsBondShare"], s["afsMoneyMktShare"] = shares(
        afsListed, afsUnlisted, afsFund, afsBond, afsMoneyMkt
    )
    allListed = [add(a, b, c) for a, b, c in zip(fvtplListed, htmListed, afsListed)]
    allUnlisted = [add(a, b, c) for a, b, c in zip(fvtplUnlisted, htmUnlisted, afsUnlisted)]
    allFund = [add(a, b, c) for a, b, c in zip(fvtplFund, htmFund, afsFund)]
    allBond = [add(a, b, c) for a, b, c in zip(fvtplBond, htmBond, afsBond)]
    allMoneyMkt = [add(a, b, c) for a, b, c in zip(fvtplMoneyMkt, htmMoneyMkt, afsMoneyMkt)]
    s["allListed"], s["allUnlisted"], s["allFund"], s["allBond"], s["allMoneyMkt"] = (
        allListed, allUnlisted, allFund, allBond, allMoneyMkt,
    )
    s["allListedShare"], s["allUnlistedShare"], s["allFundShare"], s["allBondShare"], s["allMoneyMktShare"] = shares(
        allListed, allUnlisted, allFund, allBond, allMoneyMkt
    )

    # Business-line summary (DT/LN/BLN) for the 4 segments
    s["moiGioiDT"], s["moiGioiLN"] = dtMoiGioi, lnMoiGioi
    s["moiGioiBLN"] = s["blnMoiGioi"]
    tuDoanhDT = [add(a, b, c) for a, b, c in zip(dtFvtpl, dtHtm, dtAfs)]
    tuDoanhLN = [add(a, b, c) for a, b, c in zip(lnFvtpl, lnHtm, lnAfs)]
    s["tuDoanhDT"], s["tuDoanhLN"] = tuDoanhDT, tuDoanhLN
    s["tuDoanhBLN"] = [safe_div(a, b) for a, b in zip(tuDoanhLN, tuDoanhDT)]
    s["choVayDT"], s["choVayLN"] = dtChoVay, lnChoVay
    s["choVayBLN"] = s["blnChoVay"]
    dtIB = [add(a, b) for a, b in zip(series(228), series(234))]
    lnIB = [add(a, b) for a, b in zip(series(250), series(254))]
    s["ibDT"], s["ibLN"] = dtIB, lnIB
    s["ibBLN"] = [safe_div(a, b) for a, b in zip(lnIB, dtIB)]

    lag = 4 if is_quarterly else 1
    s["moiGioiLNGrowth"] = [growth(lnMoiGioi[i], lnMoiGioi[i - lag]) if i >= lag else None for i in range(n)]
    s["tuDoanhLNGrowth"] = [growth(tuDoanhLN[i], tuDoanhLN[i - lag]) if i >= lag else None for i in range(n)]
    s["choVayLNGrowth"] = [growth(lnChoVay[i], lnChoVay[i - lag]) if i >= lag else None for i in range(n)]
    s["ibLNGrowth"] = [growth(lnIB[i], lnIB[i - lag]) if i >= lag else None for i in range(n)]

    # Margin lending
    margin, equity = series(615), series(142)
    s["margin"], s["equity"] = margin, equity
    s["marginToEquity"] = [safe_div(a, b) for a, b in zip(margin, equity)]
    marginCapacity = [sub(2 * e, m) if e is not None and m is not None else None for m, e in zip(margin, equity)]
    s["marginCapacity"] = marginCapacity
    s["marginUtilization"] = [safe_div(m, e) for m, e in zip(margin, equity)]
    s["marginGrowth"] = [growth(margin[i], margin[i - 1]) if i >= 1 else None for i in range(n)]
    marginInterest = series(224)
    annualize = 4 if is_quarterly else 1
    s["marginInterestAnnualized"] = [v * annualize if v is not None else None for v in marginInterest]
    s["marginYield"] = [safe_div(v * annualize if v is not None else None, m) for v, m in zip(marginInterest, margin)]

    # Funding cost
    shortDebt = series(94)
    s["shortDebt"] = shortDebt
    longDebt = series(121)
    interestExpense = series(265)
    avgDebt = [
        (shortDebt[i] + longDebt[i] + shortDebt[i - 1] + longDebt[i - 1]) / 2.0
        if i >= 1 and None not in (shortDebt[i], longDebt[i], shortDebt[i - 1], longDebt[i - 1])
        else None
        for i in range(n)
    ]
    s["fundingCost"] = [safe_div(-e if e is not None else None, d) for e, d in zip(interestExpense, avgDebt)]

    # NAV & investor deposits (custody/off-balance-sheet accounts)
    s["nav"] = [
        sub(add(*(r(row, i) for row in (890, 898, 903, 904, 905, 906, 907))), r(941, i)) for i in range(n)
    ]
    s["investorDeposits"] = series(907)

    # Proprietary-book investment performance (LN tự doanh / TS tự doanh)
    s["propReturn"] = [safe_div(a, b) for a, b in zip(tuDoanhLN, propTotal)]

    # Asset allocation structure
    cash = series(5)
    s["cash"], s["choVay"] = cash, f10
    s["cashShare"], s["fvtplAllocShare"], s["htmAllocShare"], s["choVayAllocShare"], s["afsAllocShare"] = shares(
        cash, tsFvtpl, tsHtm, f10, tsAfs
    )

    # Ratios needed for the "so sánh với ngành" charts — same formulas as
    # export-securities-data.py's compute_metrics (kept independent here,
    # same precedent as export-bank-detail.py not importing from
    # export-bank-data.py).
    roa, roe, epsBasic, epsDiluted, bvps = [], [], [], [], []
    grossMargin, pretaxMargin, netMargin, brokerageMargin, propProfitShare = [], [], [], [], []
    for i in range(n):
        equity_i, shares_i, totalAssets_i = r(142, i), r(176, i), r(91, i)
        revenue_i, grossProfit_i, pretax_i, netProfit_i = r(237, i), r(257, i), r(273, i), r(279, i)
        roa.append(safe_div(r(279, i), avg1(91, i)))
        roe.append(safe_div(r(280, i), avg1(142, i)))
        epsBasic.append(r(296, i) * 1e9 if r(296, i) is not None else None)
        epsDiluted.append(r(297, i) * 1e9 if r(297, i) is not None else None)
        bvps.append(safe_div(equity_i, shares_i))
        grossMargin.append(safe_div(grossProfit_i, revenue_i))
        pretaxMargin.append(safe_div(pretax_i, revenue_i))
        netMargin.append(safe_div(netProfit_i, revenue_i))
        brokerageRevenue_i, brokerageCost_i = r(227, i), r(249, i)
        brokerageMargin.append(safe_div(add(brokerageRevenue_i, brokerageCost_i), brokerageRevenue_i))
        propProfitShare.append(safe_div(add(lnFvtpl[i], lnHtm[i], lnAfs[i]), grossProfit_i))
    s["roa"], s["roe"], s["epsBasic"], s["epsDiluted"], s["bvps"] = roa, roe, epsBasic, epsDiluted, bvps
    s["grossMargin"], s["pretaxMargin"], s["netMargin"], s["brokerageMargin"], s["propProfitShare"] = (
        grossMargin, pretaxMargin, netMargin, brokerageMargin, propProfitShare,
    )

    return s


all_quarter_series = {}
all_year_series = {}
quarter_periods = None

for company_index, symbol in enumerate(COMPANIES):
    if symbol not in wb.sheetnames:
        print("MISSING SHEET:", symbol)
        continue
    ws = wb[symbol]
    sheet = Sheet(ws)
    quarter_periods = [ws.cell(row=2, column=c).value for c in QUARTER_COLS]

    all_quarter_series[symbol] = compute_series(sheet, QUARTER_COLS, True)
    all_year_series[symbol] = compute_series(sheet, YEAR_COLS, False)


def industry_avg(all_series, key, n):
    out = []
    for i in range(n):
        vals = [all_series[c][key][i] for c in COMPANIES if c in all_series and all_series[c][key][i] is not None]
        out.append(sum(vals) / len(vals) if vals else None)
    return out


INDUSTRY_KEYS = [
    "roa", "roe", "epsBasic", "epsDiluted", "bvps",
    "grossMargin", "pretaxMargin", "netMargin", "brokerageMargin", "propProfitShare",
]
quarter_industry = {k: industry_avg(all_quarter_series, k, len(quarter_periods)) for k in INDUSTRY_KEYS}
year_industry = {k: industry_avg(all_year_series, k, len(YEAR_LABELS)) for k in INDUSTRY_KEYS}

for symbol in COMPANIES:
    if symbol not in all_quarter_series:
        continue
    out = {
        "symbol": symbol,
        "quarter": {"periods": quarter_periods, "industry": quarter_industry, "company": all_quarter_series[symbol]},
        "year": {"periods": YEAR_LABELS, "industry": year_industry, "company": all_year_series[symbol]},
    }
    with open(f"{OUT_DIR}/{symbol}.json", "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
    print(symbol, "done")

print("DONE ->", OUT_DIR)
